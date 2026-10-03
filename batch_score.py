"""
Score and rank many resumes at once.

Reads an applicant sheet (.xlsx or .csv, e.g. an Airtable export), matches each
row to a PDF in a resumes folder, runs the normal scoring pipeline on it, and
writes every original column plus the scores to a CSV sorted by total score.

Usage:
    python batch_score.py applicants.xlsx resumes/ -o scores.csv
"""

import argparse
import csv
import logging
import os
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from score import score_resume, calculate_total_score

logger = logging.getLogger(__name__)

# Airtable attachment cells look like "resume.pdf (https://dl.airtable.com/...)",
# with several attachments separated by commas.
PDF_NAME_PATTERN = re.compile(r"([^,()\n]+?\.pdf)\b", re.IGNORECASE)

SCORE_COLUMNS = [
    "rank",
    "total_score",
    "max_score",
    "open_source_score",
    "self_projects_score",
    "production_score",
    "technical_skills_score",
    "bonus_points",
    "deductions",
    "key_strengths",
    "areas_for_improvement",
    "matched_resume",
    "scoring_status",
]


def read_sheet(path: Path, sheet_name: Optional[str] = None) -> Tuple[List[str], List[Dict]]:
    """Read a .xlsx/.xlsm or .csv file into (headers, rows)."""
    suffix = path.suffix.lower()

    if suffix == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = list(reader.fieldnames or [])
            rows = [dict(row) for row in reader]
        return headers, rows

    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Reading Excel files needs openpyxl: pip install openpyxl")

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        values = list(worksheet.iter_rows(values_only=True))
        workbook.close()
        if not values:
            return [], []

        headers = [
            str(h).strip() if h is not None else f"column_{i + 1}"
            for i, h in enumerate(values[0])
        ]
        rows = []
        for raw in values[1:]:
            if raw is None or all(v is None or str(v).strip() == "" for v in raw):
                continue
            rows.append(
                {h: ("" if v is None else str(v)) for h, v in zip(headers, raw)}
            )
        return headers, rows

    sys.exit(
        f"Unsupported sheet format '{suffix}'. Save it as .xlsx or .csv and try again."
    )


def normalize(text: str) -> str:
    """Lowercase and keep only letters and digits, for loose filename matching."""
    return re.sub(r"[^a-z0-9]", "", text.lower())


def index_resumes(folder: Path) -> Dict[str, Path]:
    """Map lowercased filename and normalized stem to each PDF in the folder (recursive)."""
    index = {}
    for pdf in sorted(folder.rglob("*")):
        if pdf.is_file() and pdf.suffix.lower() == ".pdf":
            index.setdefault(pdf.name.lower(), pdf)
            index.setdefault(normalize(pdf.stem), pdf)
    return index


def detect_column(headers: List[str], rows: List[Dict], kind: str) -> Optional[str]:
    """Guess the resume, name or email column from headers and cell contents."""
    lowered = {h: h.lower() for h in headers}

    if kind == "resume":
        for h in headers:
            if any(k in lowered[h] for k in ("resume", "cv")):
                return h
        # Otherwise pick the column whose cells most often mention a PDF.
        counts = {
            h: sum(1 for r in rows if ".pdf" in (r.get(h) or "").lower())
            for h in headers
        }
        best = max(counts, key=counts.get, default=None)
        return best if best and counts[best] > 0 else None

    if kind == "email":
        for h in headers:
            if "email" in lowered[h] or "e-mail" in lowered[h]:
                return h
        return None

    if kind == "name":
        preferred = ("full name", "name", "applicant name", "candidate name")
        for p in preferred:
            for h in headers:
                if lowered[h].strip() == p:
                    return h
        for h in headers:
            if "name" in lowered[h] and "user" not in lowered[h]:
                return h
        return None

    return None


def match_resume(
    row: Dict,
    index: Dict[str, Path],
    resume_column: Optional[str],
    name_column: Optional[str],
    email_column: Optional[str],
) -> Optional[Path]:
    """Find the PDF for a row: attachment filename first, then email, then name."""
    if resume_column:
        cell = row.get(resume_column) or ""
        for filename in PDF_NAME_PATTERN.findall(cell):
            filename = os.path.basename(filename.strip())
            found = index.get(filename.lower()) or index.get(
                normalize(Path(filename).stem)
            )
            if found:
                return found

    stems = {key: path for key, path in index.items() if not key.endswith(".pdf")}

    for column in (email_column, name_column):
        value = normalize(row.get(column) or "") if column else ""
        if column == email_column and value:
            # Match on the local part too, e.g. "jane.doe" from jane.doe@gmail.com.
            local = normalize((row.get(column) or "").split("@")[0])
            candidates = [value, local]
        else:
            candidates = [value]
        for candidate in candidates:
            if len(candidate) < 3:
                continue
            if candidate in stems:
                return stems[candidate]
            matches = [path for stem, path in stems.items() if candidate in stem]
            if len(set(matches)) == 1:
                return matches[0]

    return None


def evaluation_columns(evaluation) -> Dict:
    total_score, max_score = calculate_total_score(evaluation)
    scores = evaluation.scores
    return {
        "total_score": round(total_score, 2),
        "max_score": max_score,
        "open_source_score": min(scores.open_source.score, scores.open_source.max),
        "self_projects_score": min(
            scores.self_projects.score, scores.self_projects.max
        ),
        "production_score": min(scores.production.score, scores.production.max),
        "technical_skills_score": min(
            scores.technical_skills.score, scores.technical_skills.max
        ),
        "bonus_points": evaluation.bonus_points.total,
        "deductions": evaluation.deductions.total,
        "key_strengths": " | ".join(evaluation.key_strengths),
        "areas_for_improvement": " | ".join(evaluation.areas_for_improvement),
    }


def write_ranked_csv(output: Path, headers: List[str], rows: List[Dict]):
    """Sort scored rows by total score (unscored rows last) and write the CSV."""
    scored = [r for r in rows if r.get("total_score") not in (None, "")]
    unscored = [r for r in rows if r.get("total_score") in (None, "")]
    scored.sort(key=lambda r: r["total_score"], reverse=True)

    # Standard competition ranking: ties share a rank.
    previous_score, previous_rank = None, 0
    for position, row in enumerate(scored, 1):
        if row["total_score"] != previous_score:
            previous_rank, previous_score = position, row["total_score"]
        row["rank"] = previous_rank

    extra = [c for c in SCORE_COLUMNS if c not in headers]
    fieldnames = ["rank"] + [h for h in headers if h != "rank"] + [
        c for c in extra if c != "rank"
    ]

    # utf-8-sig so Excel opens names with accents correctly.
    with open(output, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(scored + unscored)


def main():
    parser = argparse.ArgumentParser(
        description="Score a sheet of applicants against a folder of resume PDFs and rank them."
    )
    parser.add_argument("sheet", type=Path, help="Applicant sheet (.xlsx or .csv)")
    parser.add_argument("resumes", type=Path, help="Folder containing resume PDFs")
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("scores.csv"), help="Output CSV path"
    )
    parser.add_argument("--sheet-name", help="Worksheet to read (default: first sheet)")
    parser.add_argument(
        "--resume-column", help="Column holding the resume attachment or filename"
    )
    parser.add_argument("--name-column", help="Column holding the applicant's name")
    parser.add_argument("--email-column", help="Column holding the applicant's email")
    args = parser.parse_args()

    if not args.sheet.exists():
        sys.exit(f"Error: sheet '{args.sheet}' does not exist.")
    if not args.resumes.is_dir():
        sys.exit(f"Error: resumes folder '{args.resumes}' does not exist.")

    headers, rows = read_sheet(args.sheet, args.sheet_name)
    if not rows:
        sys.exit("The sheet has no data rows.")

    for flag, column in (
        ("--resume-column", args.resume_column),
        ("--name-column", args.name_column),
        ("--email-column", args.email_column),
    ):
        if column and column not in headers:
            sys.exit(f"Error: {flag} '{column}' is not a column. Columns: {headers}")

    resume_column = args.resume_column or detect_column(headers, rows, "resume")
    name_column = args.name_column or detect_column(headers, rows, "name")
    email_column = args.email_column or detect_column(headers, rows, "email")
    print(
        f"Matching resumes using resume column={resume_column!r}, "
        f"name column={name_column!r}, email column={email_column!r}"
    )

    index = index_resumes(args.resumes)
    if not index:
        sys.exit(f"No PDFs found in '{args.resumes}'.")

    # Score each PDF once, even if several rows point at it.
    results: Dict[Path, Dict] = {}

    try:
        for i, row in enumerate(rows, 1):
            pdf_path = match_resume(row, index, resume_column, name_column, email_column)
            label = row.get(name_column, "") if name_column else f"row {i}"
            if pdf_path is None:
                print(f"[{i}/{len(rows)}] {label}: no matching resume found, skipping")
                row["scoring_status"] = "no resume found"
                continue

            row["matched_resume"] = pdf_path.name
            if pdf_path not in results:
                print(f"[{i}/{len(rows)}] {label}: scoring {pdf_path.name}")
                try:
                    result = score_resume(str(pdf_path))
                    if result is None or result[0] is None:
                        results[pdf_path] = {"scoring_status": "could not parse PDF"}
                    else:
                        results[pdf_path] = {
                            **evaluation_columns(result[0]),
                            "scoring_status": "scored",
                        }
                except Exception as e:
                    logger.exception(f"Failed to score {pdf_path}")
                    results[pdf_path] = {"scoring_status": f"error: {e}"}
            row.update(results[pdf_path])
    except KeyboardInterrupt:
        print("\nInterrupted, writing results scored so far...")
    finally:
        write_ranked_csv(args.output, headers, rows)

    scored = sum(1 for r in rows if r.get("scoring_status") == "scored")
    print(f"\nScored {scored}/{len(rows)} applicants. Ranked results: {args.output}")


if __name__ == "__main__":
    main()
